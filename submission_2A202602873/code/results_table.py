from __future__ import annotations

import json
import os
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    os.makedirs(results_dir, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    out_path = os.path.join(results_dir, f"{exp_id}.json")

    clean_dict = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"]
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(clean_dict, f, indent=2, ensure_ascii=False)
    return out_path


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p = Path(results_dir)
    if not p.exists():
        return []
    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            results.append(json.load(fp))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result["cfg"]
    summary = result["summary"]
    exp_id = cfg["exp_id"]

    row = dict(cfg)
    row.update(summary)
    row["figure_file"] = f"figures/{exp_id}.png"
    row["notes"] = notes

    if eval_scores is not None:
        row["eval_acc"] = eval_scores.get("accuracy", eval_scores.get("acc", eval_scores.get("eval_acc", "")))
        row["eval_macro_f1"] = eval_scores.get("macro_f1", eval_scores.get("eval_macro_f1", ""))
    else:
        row["eval_acc"] = ""
        row["eval_macro_f1"] = ""

    if isinstance(row.get("hidden"), (list, tuple)):
        row["hidden"] = "-".join(map(str, row["hidden"]))

    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path."""
    try:
        import openpyxl
    except ImportError:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
        import openpyxl

    if not os.path.exists(template_path):
        candidates = [
            Path(template_path),
            Path(__file__).resolve().parents[2] / "templates" / "experiment_table_template.xlsx",
            Path(__file__).resolve().parents[1] / "templates" / "experiment_table_template.xlsx",
            Path.cwd() / "templates" / "experiment_table_template.xlsx",
        ]
        for candidate in candidates:
            if candidate.exists():
                template_path = str(candidate)
                break

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    col_map = {cell.column: cell.value for cell in ws[1] if cell.value is not None}
    formula_cols = {
        "step0_gap_vs_lnC", "gap_val_minus_train",
        "delta_val_f1_vs_base", "beyond_noise"
    }

    start_row = 2
    for r_idx, row_data in enumerate(rows, start=start_row):
        for col_idx, col_name in col_map.items():
            if col_name in formula_cols:
                continue
            if col_name in row_data:
                ws.cell(row=r_idx, column=col_idx, value=row_data[col_name])

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
